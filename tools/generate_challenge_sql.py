import docx
import re
import os

def generate():
    doc = docx.Document('C. MPL r1 questions CHALLENGE_INDEXED_NO_ANSWERS.docx')
    
    raw_entries = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            raw_entries.append(txt)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt and txt not in raw_entries:
                    raw_entries.append(txt)

    def sanitize(text):
        text = text.replace('\u00d7', '*')  # multiplication sign
        text = text.replace('\u22c5', '*')
        text = text.replace('\u03c0', 'pi') # pi
        text = text.replace('\u00b0', ' deg')
        text = text.replace('\u2018', "'").replace('\u2019', "'")
        text = text.replace('\u201c', '"').replace('\u201d', '"')
        text = text.replace('\u2014', '-').replace('\u2013', '-')
        text = text.replace('\u2026', '...')
        
        # Subscripts: ₀₁₂₃₄₅₆₇₈₉
        subscripts = {'\u2080':'0', '\u2081':'1', '\u2082':'2', '\u2083':'3', '\u2084':'4', '\u2085':'5', '\u2086':'6', '\u2087':'7', '\u2088':'8', '\u2089':'9'}
        for k, v in subscripts.items():
            text = text.replace(k, v)

        # Superscripts: ⁰¹²³⁴⁵⁶⁷⁸⁹
        supers = {'\u2070':'0', '\u00b9':'1', '\u00b2':'2', '\u00b3':'3', '\u2074':'4', '\u2075':'5', '\u2076':'6', '\u2077':'7', '\u2078':'8', '\u2079':'9'}
        # Replace contiguous sequence of superscripts with ^<seq>
        pattern = '[' + ''.join(supers.keys()) + ']+'
        def repl_sup(m):
            s = ''.join(supers[c] for c in m.group(0))
            return '^' + s
        text = re.sub(pattern, repl_sup, text)
        
        return text.strip()

    questions = []
    for entry in raw_entries:
        match = re.search(r'c\.(\d+)\s*(?:Question:\s*)?(.*)', entry, re.DOTALL | re.IGNORECASE)
        if match:
            c_num = int(match.group(1))
            q_text = match.group(2).strip()
            q_text = re.sub(r'^Question:\s*', '', q_text, flags=re.IGNORECASE).strip()
            questions.append((c_num, q_text))

    questions.sort(key=lambda x: x[0])

    sql_lines = []
    sql_lines.append('-- ========================================================')
    sql_lines.append('-- POPULATE CHALLENGE QUESTIONS (IDs 701 - 724)')
    sql_lines.append('-- ========================================================')
    sql_lines.append('BEGIN;')
    sql_lines.append('')
    sql_lines.append('-- 1. Remove existing CHALLENGE questions and any linked sessions')
    sql_lines.append("DELETE FROM challenge_sessions WHERE question_id IN (SELECT id FROM questions WHERE type = 'CHALLENGE');")
    sql_lines.append("DELETE FROM questions WHERE type = 'CHALLENGE';")
    sql_lines.append('')
    sql_lines.append('-- 2. Insert new 24 Challenge questions')
    sql_lines.append('INSERT INTO questions (')
    sql_lines.append('    id, title, description, test_cases, type, difficulty, reward_value, sub_type, starter_code, compare_mode, order_index, cpu_time_limit, wall_time_limit, memory_limit_kb, points')
    sql_lines.append(') VALUES')

    val_rows = []
    for c_num, qtext in questions:
        q_id = 700 + c_num
        title = f'Challenge: c.{c_num}'
        clean_desc = sanitize(qtext)
        desc = f"{clean_desc}\n\nType: Challenge Round | 1v1 Battle"
        
        row = f"({q_id}, $t${title}$t$, $d${desc}$d$, '[]', 'CHALLENGE', 'HARD', 100, NULL, NULL, 'TRIM', 0, 5.0, 10.0, 256000, 100)"
        val_rows.append(row)

    sql_lines.append(',\n'.join(val_rows) + ';')
    sql_lines.append('')
    sql_lines.append('-- 3. Reset ID sequence')
    sql_lines.append("SELECT setval('questions_id_seq', (SELECT COALESCE(MAX(id), 1) FROM questions));")
    sql_lines.append('')
    sql_lines.append('COMMIT;')

    output_path = 'populate_challenge_questions.sql'
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(sql_lines))

    print(f"Generated {output_path} with {len(val_rows)} questions successfully.")

if __name__ == '__main__':
    generate()
